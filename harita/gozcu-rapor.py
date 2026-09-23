#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ÜSTAD GÖZCÜ — GÜNLÜK RAPOR (renkli Word)
========================================
Panel sunucusundan (port 8811) o günün durumunu ve rapor metnini alır,
Kenan'ın istediği biçimde RENKLİ bir .docx üretir ve Masaüstü'ne bırakır.

Kullanım:
    python gozcu-rapor.py            → Masaüstü\\USTAD-GOZCU-RAPOR-<tarih>.docx
    python gozcu-rapor.py --kopya    → ayrıca panele de kaydeder

Sunucu kapalıysa: rapor yine üretilir ama "sunucu kapalı" notu düşülür.
"""
import datetime
import json
import os
import sys
import urllib.error
import urllib.request

KOK = "http://127.0.0.1:8811"
MASAUSTU = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop")
if not os.path.isdir(MASAUSTU):
    MASAUSTU = os.path.join(os.path.expanduser("~"), "Desktop")

# renkler (Kenan'ın panel paleti)
TEAL = "2BF0C8"
AMBER = "FFD23B"
RED = "FF3B4D"
BLUE = "3D8BFF"
GRI = "6B7C86"
KOYU = "0A1014"


def al(yol, sure=45):
    with urllib.request.urlopen(KOK + yol, timeout=sure) as c:
        return json.loads(c.read().decode("utf-8", "replace"))


def veri_toplan():
    d = {"rapor": None, "durum": None, "hava": None, "deprem": None,
         "uyari": None, "piyasa": None, "hata": []}
    for anahtar, yol in (("rapor", "/veri/rapor"), ("durum", "/veri/durum"),
                         ("hava", "/veri/hava"), ("deprem", "/veri/deprem"),
                         ("uyari", "/veri/uyari"), ("piyasa", "/veri/piyasa")):
        try:
            d[anahtar] = al(yol)
        except Exception as hata:
            d["hata"].append("%s: %s" % (yol, str(hata)[:60]))
    return d


def satir_serit(doc, metin, renk):
    """Renkli başlık şeridi (Word tablosu ile)."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    hucre = t.rows[0].cells[0]
    hucre.text = ""
    p = hucre.paragraphs[0]
    kosu = p.add_run("  " + metin)
    kosu.bold = True
    kosu.font.size = None  # varsayılan
    from docx.shared import Pt, RGBColor
    kosu.font.size = Pt(13)
    kosu.font.color.rgb = RGBColor.from_string("FFFFFF")
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    golge = OxmlElement("w:shd")
    golge.set(qn("w:fill"), renk)
    hucre._tc.get_or_add_tcPr().append(golge)
    doc.add_paragraph()
    return t


def tablo(doc, basliklar, satirlar):
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    from docx.shared import Pt
    for i, b in enumerate(basliklar):
        h = t.rows[0].cells[i]
        h.text = ""
        kosu = h.paragraphs[0].add_run(b)
        kosu.bold = True
        kosu.font.size = Pt(10)
    for s in satirlar:
        hucreler = t.add_row().cells
        for i, x in enumerate(s):
            hucreler[i].text = ""
            hucreler[i].paragraphs[0].add_run(str(x)).font.size = Pt(10)
    return t


def uret(kopya=False):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = veri_toplan()
    simdi = datetime.datetime.now()
    gun = simdi.strftime("%d.%m.%Y")
    doc = Document()
    genislik = doc.sections[0]

    # ---- kapak ----
    baslik = doc.add_paragraph()
    baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER
    k = baslik.add_run("ÜSTAD GÖZCÜ")
    k.bold = True
    k.font.size = Pt(30)
    k.font.color.rgb = RGBColor.from_string(TEAL)
    alt = doc.add_paragraph()
    alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a = alt.add_run("GÜNLÜK DÜNYA İZLEME RAPORU")
    a.bold = True
    a.font.size = Pt(13)
    a.font.color.rgb = RGBColor.from_string(AMBER)
    tarih = doc.add_paragraph()
    tarih.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = tarih.add_run(gun + " · " + simdi.strftime("%H:%M") + " · ÜSTAD KENAN KUZUCU")
    t.font.size = Pt(11)
    t.font.color.rgb = RGBColor.from_string(GRI)
    doc.add_paragraph()

    if d["hata"]:
        u = doc.add_paragraph()
        uu = u.add_run("! Kaynak uyarısı: " + " | ".join(d["hata"][:3]))
        uu.font.size = Pt(9)
        uu.font.color.rgb = RGBColor.from_string(RED)

    # ---- panel rapor metni + katman sayıları ----
    rp = d.get("rapor") or {}
    satir_serit(doc, "1. CANLI DURUM (panel raporu)", TEAL)
    if rp.get("metin"):
        sayac = 0
        for satir in str(rp["metin"]).split("\n"):
            s = satir.strip()
            if not s or len(s) < 4:
                continue
            if not any(ch.isdigit() for ch in s):
                continue
            p = doc.add_paragraph(s)
            for kosu in p.runs:
                kosu.font.size = Pt(10)
            sayac += 1
            if sayac >= 16:
                break
    else:
        p = doc.add_paragraph("Panel raporu alınamadı (sunucu kapalı olabilir).")
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.color.rgb = RGBColor.from_string(GRI)
    doc.add_paragraph()

    # ---- hava durumu ----
    hava = d.get("hava") or {}
    iller = hava.get("iller") or []
    if iller:
        satir_serit(doc, "2. TÜRKİYE HAVA DURUMU (" + str(len(iller)) + " il)", BLUE)
        sirali = sorted([x for x in iller if x.get("t") is not None],
                        key=lambda x: x.get("t") or 0)
        en_soguk = sirali[:5]
        en_sicak = sirali[-5:][::-1]
        tablo(doc, ["En serin il", "Sıcaklık"], [[x["ad"], "%.1f °C" % x["t"]] for x in en_soguk])
        doc.add_paragraph()
        tablo(doc, ["En sıcak il", "Sıcaklık"], [[x["ad"], "%.1f °C" % x["t"]] for x in en_sicak])
        yagisli = [x for x in iller if (x.get("y") or 0) > 0]
        p = doc.add_paragraph("Yağış bildiren il sayısı: %d" % len(yagisli))
        p.runs[0].font.size = Pt(10)
        doc.add_paragraph()

    # ---- depremler ----
    dep = d.get("deprem") or {}
    dep_liste = dep.get("depremler") or dep.get("liste") or []
    if dep_liste:
        satir_serit(doc, "3. SON 24 SAATİN DEPREMLERİ", AMBER)
        try:
            sirali = sorted(dep_liste, key=lambda x: float(x.get("m") or x.get("buyukluk") or 0),
                            reverse=True)[:10]
        except Exception:
            sirali = dep_liste[:10]
        satirlar = []
        for x in sirali:
            satirlar.append([str(x.get("yer") or x.get("ad") or "—")[:44],
                             str(x.get("m") or x.get("buyukluk") or "—")])
        tablo(doc, ["Yer", "Büyüklük"], satirlar)
        doc.add_paragraph()

    # ---- uyarılar ----
    uy = d.get("uyari") or {}
    bolgeler = uy.get("bolgeler") or []
    if bolgeler:
        satir_serit(doc, "4. UYARI BÖLGELERİ (coğrafi çit)", RED)
        satirlar = []
        for b in bolgeler:
            olaylar = b.get("olaylar") or []
            satirlar.append([b.get("ad") or "—",
                             ", ".join(b.get("katlar") or []) or "—",
                             str(len(olaylar)) + " olay"])
        tablo(doc, ["Bölge", "İzlenen katmanlar", "Yakalanan"], satirlar)
        doc.add_paragraph()

    # ---- piyasa ----
    piy = d.get("piyasa") or {}
    piy_liste = piy.get("piyasalar") or piy.get("liste") or []
    if piy_liste:
        satir_serit(doc, "5. PİYASALAR (ilk 12)", TEAL)
        satirlar = []
        for x in piy_liste[:12]:
            satirlar.append([str(x.get("ad") or x.get("sembol") or "—"),
                             str(x.get("fiyat") or x.get("son") or "—"),
                             str(x.get("degisim") or x.get("yuzde") or "—")])
        tablo(doc, ["Enstrüman", "Fiyat", "Değişim %"], satirlar)
        doc.add_paragraph()

    # ---- footer ----
    doc.add_paragraph()
    son = doc.add_paragraph()
    son.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s = son.add_run("ÜSTAD GÖZCÜ · açık veriyle çalışan yerel izleme paneli · "
                    "ÜSTAD KENAN KUZUCU")
    s.font.size = Pt(9)
    s.font.color.rgb = RGBColor.from_string(GRI)

    ad = "USTAD-GOZCU-RAPOR-" + simdi.strftime("%Y-%m-%d") + ".docx"
    tam = os.path.join(MASAUSTU, ad)
    doc.save(tam)
    if kopya:
        try:
            kopya_yol = os.path.join("C:\\Users\\kenan\\OneDrive\\Desktop\\USTAD-GOZCU",
                                     "raporlar")
            os.makedirs(kopya_yol, exist_ok=True)
            doc.save(os.path.join(kopya_yol, ad))
        except Exception:
            pass
    return tam


if __name__ == "__main__":
    try:
        yol = uret("--kopya" in sys.argv)
        print("Rapor hazır:", yol)
    except Exception as hata:
        print("RAPOR ÜRETİLEMEDİ: %s: %s" % (type(hata).__name__, hata))
        sys.exit(1)
